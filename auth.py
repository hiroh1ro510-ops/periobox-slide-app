import io
import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import streamlit as st
from auth import get_secret_value, load_local_env, require_authorized_user
from openai import OpenAI, OpenAIError
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pypdf import PdfReader


load_local_env()
APP_TITLE = "PerioBox 4枚スライド自動生成 MVP"
DEFAULT_MODEL = get_secret_value("OPENAI_MODEL", "gpt-4o-mini")
MAX_INPUT_CHARS = 60000
OUTPUT_DIR = Path("outputs")
TOOTH_ICON_PATH = Path("assets/periobox_tooth.png")
BACKGROUND_IMAGE_PATH = Path("assets/periobox_background.jpeg")
CQ_TRIANGLE_PATH = Path("assets/cq_answer_triangle.png")
CQ_UNDERLINE_PATH = Path("assets/cq_answer_underline.png")
DOCTOR_COMMENT_BG_PATH = Path("assets/doctor_comment_bg.png")
DOCTOR_COMMENT_ICON_PATH = Path("assets/doctor_comment_icon.png")
SELECTION_POINT_UNDERLINE_PATH = Path("assets/selection_point_underline.png")
FONT_NAME = "Century Gothic"
JP_FONT_NAME = "MS PGothic"
TEXT_COLOR = RGBColor(37, 37, 37)
MUTED_COLOR = RGBColor(115, 115, 115)
QUESTION_FILL = RGBColor(232, 243, 225)
ACCENT_PURPLE = RGBColor(142, 130, 198)
HEADER_BLUE = RGBColor(203, 239, 251)
HEADER_LINE_BLUE = RGBColor(190, 236, 250)


def extract_pdf_text(pdf_file: io.BytesIO) -> str:
    reader = PdfReader(pdf_file)
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = re.sub(r"\s+", " ", text).strip()
        if text:
            pages.append(f"[Page {index}]\n{text}")
    return "\n\n".join(pages)


def trim_text(text: str, max_chars: int = MAX_INPUT_CHARS) -> str:
    if len(text) <= max_chars:
        return text
    head = text[: int(max_chars * 0.7)]
    tail = text[-int(max_chars * 0.3) :]
    return f"{head}\n\n[...中略: 入力が長いため冒頭と末尾を使用...]\n\n{tail}"


def get_openai_client(api_key: str) -> OpenAI:
    api_key = api_key.strip() or get_secret_value("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY がサーバー側に設定されていません。管理者に確認してください。")
    return OpenAI(api_key=api_key)


def call_openai_json(client: OpenAI, model: str, system: str, prompt: str) -> Dict[str, Any]:
    try:
        response = client.chat.completions.create(
            model=model,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": prompt},
            ],
            temperature=0.1,
        )
        content = response.choices[0].message.content or "{}"
    except OpenAIError as error:
        message = str(error)
        if "insufficient_quota" in message or "exceeded your current quota" in message:
            raise RuntimeError(
                "OpenAI APIの利用枠または残高が不足しています。"
                "OpenAI PlatformのBilling/Usageで課金設定、残高、利用上限を確認してください。"
            ) from error
        raise RuntimeError(f"OpenAI APIエラー: {message}") from error

    try:
        return json.loads(content)
    except json.JSONDecodeError as error:
        raise RuntimeError("OpenAI APIのJSON出力を読み取れませんでした。もう一度生成してください。") from error


def call_openai_text(client: OpenAI, model: str, system: str, prompt: str) -> str:
    try:
        response = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": prompt},
            ],
            temperature=0.1,
        )
        return response.choices[0].message.content or ""
    except OpenAIError as error:
        message = str(error)
        if "insufficient_quota" in message or "exceeded your current quota" in message:
            raise RuntimeError(
                "OpenAI APIの利用枠または残高が不足しています。"
                "OpenAI PlatformのBilling/Usageで課金設定、残高、利用上限を確認してください。"
            ) from error
        raise RuntimeError(f"OpenAI APIエラー: {message}") from error


def build_section_extraction_prompt(pdf_text: str) -> str:
    return f"""
以下の歯科医療論文PDF本文を、論文セクションごとに整理してください。

目的:
- 後続ステップで忠実な日本語要約を作るため、本文構造をできるだけ正確に復元する。
- この段階ではPerioBox風の言い換えをしない。

厳守事項:
- 出力はJSONのみ。
- 本文にない内容は推測しない。
- セクションが本文から見つからない場合は「本文から判別困難」とする。
- Introduction/Backgroundは特に丁寧に抽出する。
- Methods/Resultsでは対象、介入、比較、評価項目、数値、P値、CIを残す。
- Discussion/Conclusionでは著者の解釈と限界を分ける。

JSONスキーマ:
{{
  "paper_metadata": {{
    "title": "論文タイトル",
    "authors": "著者",
    "journal": "雑誌名",
    "year": "年",
    "article_type": "論文タイプ"
  }},
  "sections": {{
    "abstract": "Abstractの要点",
    "introduction": "Introduction/Backgroundの忠実な抽出要約",
    "methods": "Methodsの忠実な抽出要約",
    "results": "Resultsの忠実な抽出要約",
    "discussion": "Discussionの忠実な抽出要約",
    "conclusion": "Conclusionの忠実な抽出要約",
    "limitations": "限界",
    "figures_tables": "図表・表の内容"
  }},
  "key_data": {{
    "participants": "対象",
    "interventions_or_exposures": ["介入・曝露・群1", "群2"],
    "comparisons": ["比較1", "比較2"],
    "outcomes": ["評価項目1", "評価項目2"],
    "numeric_results": ["数値結果1", "数値結果2"],
    "statistical_notes": ["統計情報1", "統計情報2"]
  }},
  "extraction_warnings": ["PDF抽出上の注意点"]
}}

PDF本文:
{trim_text(pdf_text, 90000)}
""".strip()


def build_faithful_summary_prompt(sections: Dict[str, Any]) -> str:
    return f"""
以下は歯科医療論文PDFからセクション別に抽出した内容です。
これをもとに、原文に忠実な日本語の構造化要約を作成してください。

目的:
- PerioBoxスライド化の前段階として、翻訳精度を上げる。
- 原文の意味、対象、方法、結果、限界、数値を落とさない。

厳守事項:
- 出力はJSONのみ。
- この段階ではスライド用に過度に短縮しない。
- 不明点は「本文から判別困難」と書く。
- introduction_jaはIntroduction/Backgroundのみをもとに約400字で書く。
- methods_ja/results_ja/discussion_jaは原文に忠実に、臨床的に重要な情報を残す。
- 著者の結論と、あなたの臨床解釈を混ぜない。

JSONスキーマ:
{{
  "paper": {{
    "title_en": "英語タイトル",
    "citation_short": "筆頭著者 et al. ジャーナル名. 年",
    "article_type": "論文タイプ"
  }},
  "faithful_summary": {{
    "introduction_ja": "Introductionを約400字で忠実に要約",
    "clinical_question_source": "本文から導ける臨床疑問",
    "methods_ja": "Methodsの忠実な日本語要約",
    "participants_ja": "対象",
    "interventions_or_groups_ja": ["群・介入・曝露1", "群・介入・曝露2"],
    "outcomes_ja": ["評価項目1", "評価項目2"],
    "results_ja": "Resultsの忠実な日本語要約",
    "key_numeric_data_ja": ["数値1", "数値2", "数値3"],
    "statistical_notes_ja": ["統計1", "統計2"],
    "discussion_ja": "Discussionの忠実な日本語要約",
    "limitations_ja": ["限界1", "限界2"],
    "figures_tables_ja": "図表の忠実な日本語要約"
  }}
}}

セクション別抽出JSON:
{json.dumps(sections, ensure_ascii=False)}
""".strip()


def build_slide_prompt(faithful_summary: Dict[str, Any]) -> str:
    return f"""
あなたは歯周病学・インプラント周囲疾患に詳しい歯科医師向け論文解説者です。
以下の「原文に忠実な日本語構造化要約」から、実際のPerioBox抄読スライドに近い4枚構成PowerPoint用の日本語要約を作成してください。

厳守事項:
- 出力はJSONのみ。Markdownや説明文を付けない。
- 不明な情報は推測せず「本文から判別困難」と書く。
- 元の忠実要約から外れた新情報を追加しない。
- 歯科医師向けの抄読スライドとして、短すぎず、臨床背景・方法・結果・解釈が伝わる密度にする。
- 文体はPerioBox風に、断定しすぎず「〜が示された」「〜の参考になる」「〜可能性がある」を適切に使う。
- 数値、P値、信頼区間、効果量が本文にあれば優先して拾う。
- clinical_question_jpは「〜か？」で終わる日本語の問いにする。
- key_message_jpは問いに対する一文の答えにする。
- background_overviewはPDF本文のIntroduction/Background/緒言に相当する箇所を優先して要約する。
- background_overviewは日本語で約400字（380〜430字目安）にする。
- background_overviewには研究方法・結果・結論を混ぜすぎず、「臨床上の課題」「既存知見」「未解決点」「本研究の目的」が自然につながるように書く。
- article_typeはレビュー論文なら Systematic Review / Meta-Analysis / Literature Review / Narrative Review など、レビュー以外なら RCT / Retrospective study / Case report / Experimental study などと明確に書く。
- selection_pointsは「①」「②」を付けず、短い2項目にする。
- doctor_commentは臨床家が読む自然な長めのコメントにする。限界にも触れ、安易な臨床応用を避ける。
- 研究論文のslide2.design_summaryは「比較研究」などの曖昧な説明ではなく、論文内で言及されている可能性が高い一般的な研究デザイン名にする。例: ランダム化比較試験、後ろ向きコホート研究、前向きコホート研究、横断研究、症例対照研究、in vitro比較研究。
- slide2.outcomesは、論文内で使われている略称があれば必ず日本語名の後ろに括弧で付ける。例: プロービングデプス（PD）、臨床的アタッチメントレベル（CAL）、辺縁骨レベル（MBL）、出血（BOP）。
- 研究論文のslide2.design_schemaは、2枚目中段にPowerPointの編集可能な図形（角丸四角、矢印、テキスト）で組める範囲の図・表として再構成できる形で書く。
- 比較試験では、何と何を比較しているか、対象から評価までの流れ、群分け、評価タイミングが視覚的に分かるようにする。
- slide2.design_schema.stepsは中段の流れ、groupsは比較群、labelsは図中ラベルに使える短語にする。
- 研究論文のslide3.main_resultsは、統計的有意差がある結果を優先してメインに記載する。
- slide3.main_resultsの各項目には、有意差があったかどうかを本文の表現に沿って必ず含める。例: 「〜は有意に改善した（p<0.05）」「〜に有意差は認められなかった」。
- 有意差が本文から判別困難な場合は「有意差は本文から判別困難」と添える。
- 研究論文のslide3.evidence_schemaは、Results内の主要結果に関連するFigure/Tableの内容を参考にしつつ、著作物を複製せず、PowerPointの編集可能な図形・表・テキストで作れる範囲の根拠表として再構成できる形で書く。
- evidence_schema.itemsは、比較軸・評価項目・数値・統計・臨床的解釈が伝わる3〜4項目に絞る。

JSONスキーマ:
{{
  "paper": {{
    "title_en": "英語論文タイトル",
    "citation": "筆頭著者 et al. ジャーナル名. 年",
    "article_type": "RCT / Systematic Review / Retrospective study / Case report 等"
  }},
  "slide1": {{
    "clinical_question_jp": "日本語の問い",
    "key_message_jp": "問いへの結論を1文で",
    "background_heading": "論文の背景と概要",
    "background_overview": "Introductionを中心に、臨床上の課題・既存知見・未解決点・本研究目的を約400字で要約"
  }},
  "slide2": {{
    "design_summary": "一般的な研究デザイン名。例: ランダム化比較試験、後ろ向きコホート研究、in vitro比較研究",
    "participants": "対象患者・対象資料",
    "selection_criteria": ["選定条件1", "選定条件2", "除外条件1"],
    "methods": ["方法1", "方法2", "方法3"],
    "groups_or_comparison": ["群分け・比較1", "群分け・比較2"],
    "outcomes": ["評価項目1（略称）", "評価項目2（略称）", "評価項目3（略称）"],
    "design_schema": {{
      "title": "中段図表の短いタイトル",
      "steps": ["対象/処置の流れ1", "流れ2", "流れ3"],
      "groups": [
        {{"name": "群・比較対象1", "description": "短い説明"}},
        {{"name": "群・比較対象2", "description": "短い説明"}}
      ],
      "labels": ["図中ラベル1", "図中ラベル2"]
    }}
  }},
  "slide3": {{
    "main_results": ["統計的有意差の有無を含む主な結果1", "統計的有意差の有無を含む主な結果2", "統計的有意差の有無を含む主な結果3"],
    "key_numbers": ["重要な数値1", "重要な数値2", "重要な数値3"],
    "statistical_notes": ["統計的有意差・P値1", "統計的有意差・P値2"],
    "evidence_schema": {{
      "title": "主要結果の根拠を示す短いタイトル",
      "items": [
        {{"label": "比較軸・評価項目", "value": "数値または所見", "interpretation": "臨床的な読み取り"}},
        {{"label": "比較軸・評価項目", "value": "数値または所見", "interpretation": "臨床的な読み取り"}}
      ]
    }},
    "more_detail": "さらに詳しく：図表や補足結果の要約を1〜2文"
  }},
  "slide4": {{
    "selection_points": ["論文選定のポイント1", "論文選定のポイント2"],
    "doctor_comment": "抄読ドクターからひとこと。臨床的意義、使いどころ、注意点、限界を含めて5〜8文",
    "limitations": ["限界1", "限界2"]
  }}
}}

原文に忠実な日本語構造化要約JSON:
{json.dumps(faithful_summary, ensure_ascii=False)}
""".strip()


def extract_sections_with_openai(client: OpenAI, model: str, pdf_text: str) -> Dict[str, Any]:
    return call_openai_json(
        client,
        model,
        "あなたは医学論文PDFの崩れた抽出テキストから、論文セクションを正確に復元する専門家です。",
        build_section_extraction_prompt(pdf_text),
    )


def create_faithful_summary_with_openai(client: OpenAI, model: str, sections: Dict[str, Any]) -> Dict[str, Any]:
    return call_openai_json(
        client,
        model,
        "あなたは歯科医療論文を原文に忠実に日本語へ構造化要約する専門家です。",
        build_faithful_summary_prompt(sections),
    )


def create_slide_summary_with_openai(client: OpenAI, model: str, faithful_summary: Dict[str, Any]) -> Dict[str, Any]:
    return call_openai_json(
        client,
        model,
        "あなたは歯科医師向けPerioBox抄読スライドを作る専門家です。",
        build_slide_prompt(faithful_summary),
    )


def as_list(value: Any) -> List[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item) for item in value if str(item).strip()]
    return [str(value)]


def get_nested(data: Dict[str, Any], *keys: str, default: str = "本文から判別困難") -> str:
    current: Any = data
    for key in keys:
        if not isinstance(current, dict) or key not in current:
            return default
        current = current[key]
    if current is None or str(current).strip() == "":
        return default
    return str(current)


def font_size_for(text: str, base: int, long: int, threshold: int) -> int:
    return long if len(str(text)) >= threshold else base


def disable_shadow(shape):
    if hasattr(shape, "shadow"):
        shape.shadow.inherit = False
    return shape


def short_citation(citation: str) -> str:
    citation = re.sub(r"https?://\S+", "", str(citation))
    citation = re.sub(r"\bdoi\s*:?\s*\S+", "", citation, flags=re.IGNORECASE)
    citation = re.sub(r"\s+", " ", citation).strip(" .")
    if not citation:
        return "本文から判別困難"

    year_match = re.search(r"\b(19|20)\d{2}\b", citation)
    year = year_match.group(0) if year_match else ""
    before_year = citation[: year_match.start()].strip(" .;,") if year_match else citation

    author_part = before_year
    journal_part = ""
    if "." in before_year:
        parts = [part.strip(" .;,") for part in before_year.split(".") if part.strip(" .;,")]
        if len(parts) >= 2:
            author_part = parts[0]
            for part in parts[1:]:
                if re.fullmatch(r"et\s+al", part, flags=re.IGNORECASE):
                    continue
                if re.search(r"[A-Za-z]", part) and not re.match(r"^\d", part):
                    journal_part = part
                    break
    elif ";" in before_year:
        parts = [part.strip(" .;,") for part in before_year.split(";") if part.strip(" .;,")]
        author_part = parts[0]
        journal_part = parts[-1] if len(parts) > 1 else ""

    first_author = author_part
    first_author = re.sub(r"\bet\s+al\b\.?", "", first_author, flags=re.IGNORECASE).strip(" .;,")
    if "," in first_author:
        name_parts = [part.strip() for part in first_author.split(",") if part.strip()]
        first_author = " ".join(name_parts[:2]) if len(name_parts) >= 2 and len(name_parts[1]) <= 8 else name_parts[0]
    elif " and " in first_author:
        first_author = first_author.split(" and ")[0].strip()

    if not journal_part:
        known_journals = [
            "J Periodontol",
            "J Clin Periodontol",
            "Clin Oral Investig",
            "Clin Implant Dent Relat Res",
            "Clin Adv Periodontics",
        ]
        for journal in known_journals:
            if journal.lower() in citation.lower():
                journal_part = journal
                break

    pieces = [first_author]
    if "et al" not in first_author.lower():
        pieces.append("et al.")
    if journal_part:
        pieces.append(f"{journal_part}.")
    if year:
        pieces.append(year)
    return " ".join(piece for piece in pieces if piece).strip(" .") + "."


def is_review_article(summary: Dict[str, Any]) -> bool:
    article_type = get_nested(summary, "paper", "article_type", default="")
    title = get_nested(summary, "paper", "title_en", default="")
    combined = f"{article_type} {title}".lower()
    review_terms = [
        "systematic review",
        "meta-analysis",
        "meta analysis",
        "literature review",
        "narrative review",
        "scoping review",
        "umbrella review",
        "review",
    ]
    non_review_terms = [
        "retrospective",
        "prospective",
        "randomized",
        "randomised",
        "trial",
        "case report",
        "cohort",
        "cross-sectional",
        "experimental",
        "in vitro",
        "in vivo",
    ]
    return any(term in combined for term in review_terms) and not any(term in combined for term in non_review_terms)


def add_textbox(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    font_size: int = 18,
    bold: bool = False,
    color: RGBColor = TEXT_COLOR,
    align=PP_ALIGN.LEFT,
    font_name: str = FONT_NAME,
    margin: float = 0.02,
    rotation: float = 0,
    word_wrap: bool = True,
    fill_color: Optional[RGBColor] = None,
):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    disable_shadow(shape)
    if fill_color is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        shape.line.fill.background()
    shape.rotation = rotation
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = word_wrap
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_multiline_text(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    font_size: int = 16,
    bold: bool = False,
    color: RGBColor = TEXT_COLOR,
    line_spacing: float = 1.15,
    font_name: str = FONT_NAME,
):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    disable_shadow(shape)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    lines = [line.strip() for line in str(text).splitlines() if line.strip()] or [str(text)]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = line_spacing
        run = paragraph.add_run()
        run.text = line
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def add_slide_background(slide):
    if BACKGROUND_IMAGE_PATH.exists():
        disable_shadow(slide.shapes.add_picture(str(BACKGROUND_IMAGE_PATH), 0, 0, width=Inches(10), height=Inches(7.5)))


def add_header_number(slide, box_number: str):
    add_textbox(slide, box_number or "000", 0.25, 0.38, 1.15, 0.55, font_size=30, color=RGBColor(130, 130, 130), margin=0)
    add_textbox(slide, "Perio Box", 0.42, 1.12, 0.8, 0.2, font_size=8, bold=True, color=RGBColor(60, 60, 60), margin=0)


def add_title_header(slide, box_number: str, title: str, citation: str):
    box_left = 0.26
    box_top = 0.31
    box_width = 1.2
    box_height = 1.25
    number_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box_left), Inches(box_top), Inches(box_width), Inches(box_height))
    disable_shadow(number_box)
    number_box.fill.solid()
    number_box.fill.fore_color.rgb = HEADER_BLUE
    number_box.line.color.rgb = HEADER_BLUE
    if TOOTH_ICON_PATH.exists():
        logo_height = 1.12
        logo_width = logo_height * 102 / 160
        logo_left = box_left + (box_width - logo_width) / 2
        logo_top = box_top + (box_height - logo_height) / 2
        disable_shadow(slide.shapes.add_picture(str(TOOTH_ICON_PATH), Inches(logo_left), Inches(logo_top), height=Inches(logo_height)))

    add_textbox(
        slide,
        box_number or "000",
        0.25,
        0.58,
        1.22,
        0.62,
        font_size=48,
        color=RGBColor(115, 115, 115),
        align=PP_ALIGN.CENTER,
        font_name=FONT_NAME,
        margin=0,
        word_wrap=False,
    )

    header_line_height = 0.035
    header_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(box_left + box_width),
        Inches(box_top + box_height - header_line_height),
        Inches(8.25),
        Inches(header_line_height),
    )
    disable_shadow(header_line)
    header_line.fill.solid()
    header_line.fill.fore_color.rgb = HEADER_LINE_BLUE
    header_line.line.color.rgb = HEADER_LINE_BLUE

    add_textbox(
        slide,
        title,
        1.65,
        0.42,
        7.65,
        0.86,
        font_size=18,
        bold=True,
        color=RGBColor(0, 0, 0),
        font_name=FONT_NAME,
        margin=0,
    )
    add_textbox(
        slide,
        short_citation(citation),
        5.15,
        1.28,
        4.15,
        0.28,
        font_size=12,
        bold=False,
        color=RGBColor(0, 0, 0),
        align=PP_ALIGN.RIGHT,
        font_name=FONT_NAME,
        margin=0,
    )


def add_section_heading(slide, text: str, left: float, top: float, width: float = 3.0):
    add_textbox(slide, text, left, top, width, 0.42, font_size=20, bold=True, color=TEXT_COLOR, margin=0)


def add_bullet_lines(slide, items: List[str], left: float, top: float, width: float, height: float, font_size: int = 16):
    text = "\n".join(f"・{item}" for item in (items or ["本文から判別困難"]))
    add_multiline_text(slide, text, left, top, width, height, font_size=font_size)


def add_pill(slide, text: str, left: float, top: float, width: float, height: float, fill: RGBColor, font_size: int = 20, bold: bool = True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    disable_shadow(shape)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.2)
    frame.margin_right = Inches(0.15)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT_NAME
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    return shape


def add_result_card(slide, title: str, body: str, left: float, top: float, width: float, height: float):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    disable_shadow(shape)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(248, 248, 248)
    shape.line.color.rgb = RGBColor(205, 205, 205)
    add_textbox(slide, title, left + 0.14, top + 0.12, width - 0.28, 0.28, font_size=13, bold=True, color=RGBColor(75, 75, 75))
    add_multiline_text(slide, body, left + 0.18, top + 0.5, width - 0.36, height - 0.62, font_size=15)


def set_shape_text(
    shape,
    text: str,
    font_size: int = 16,
    bold: bool = False,
    color: RGBColor = TEXT_COLOR,
    font_name: str = JP_FONT_NAME,
    align=PP_ALIGN.LEFT,
    margin_left: float = 0.12,
    margin_right: float = 0.12,
    margin_top: float = 0.06,
    margin_bottom: float = 0.04,
    line_spacing: float = 1.05,
):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin_left)
    frame.margin_right = Inches(margin_right)
    frame.margin_top = Inches(margin_top)
    frame.margin_bottom = Inches(margin_bottom)
    lines = [line.strip() for line in str(text).splitlines() if line.strip()] or ["本文から判別困難"]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        run = paragraph.add_run()
        run.text = line
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def add_rounded_box(slide, left: float, top: float, width: float, height: float, fill: Optional[RGBColor], line: RGBColor, line_width: float = 1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    disable_shadow(shape)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_arrow(slide, left: float, top: float, width: float, height: float):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(width), Inches(height))
    disable_shadow(arrow)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(214, 223, 238)
    arrow.line.color.rgb = RGBColor(104, 116, 132)
    arrow.line.width = Pt(0.8)
    return arrow


def parse_page_number(value: Any) -> Optional[int]:
    if value is None:
        return None
    match = re.search(r"\d+", str(value))
    if not match:
        return None
    page_number = int(match.group(0))
    return page_number if page_number > 0 else None


def get_figure_table_reference(summary: Dict[str, Any], slide_key: str) -> Dict[str, str]:
    reference = summary.get(slide_key, {}).get("figure_table_reference")
    if not isinstance(reference, dict):
        return {
            "page_number": "",
            "label": "Figure/Table",
            "caption_or_content": "本文から判別困難",
            "reason": "本文から判別困難",
        }
    return {
        "page_number": str(reference.get("page_number", "") or ""),
        "label": str(reference.get("label", "") or "Figure/Table").strip(),
        "caption_or_content": str(reference.get("caption_or_content", "") or "本文から判別困難").strip(),
        "reason": str(reference.get("reason", "") or "本文から判別困難").strip(),
    }


def normalize_figure_table_reference(reference: Any) -> Dict[str, str]:
    if not isinstance(reference, dict):
        return {
            "page_number": "",
            "label": "Figure/Table",
            "caption_or_content": "本文から判別困難",
            "reason": "本文から判別困難",
        }
    return {
        "page_number": str(reference.get("page_number", "") or ""),
        "label": str(reference.get("label", "") or "Figure/Table").strip(),
        "caption_or_content": str(reference.get("caption_or_content", "") or "本文から判別困難").strip(),
        "reason": str(reference.get("reason", "") or "本文から判別困難").strip(),
    }


def get_figure_table_candidates(summary: Dict[str, Any], slide_key: str) -> List[Dict[str, str]]:
    slide_data = summary.get(slide_key, {})
    raw_candidates = slide_data.get("figure_table_candidates") if isinstance(slide_data, dict) else None
    candidates: List[Dict[str, str]] = []
    if isinstance(raw_candidates, list):
        for item in raw_candidates:
            candidates.append(normalize_figure_table_reference(item))
    if not candidates:
        candidates.append(get_figure_table_reference(summary, slide_key))
    while len(candidates) < 3:
        candidates.append(
            {
                "page_number": "",
                "label": f"候補{len(candidates) + 1}",
                "caption_or_content": "本文から判別困難",
                "reason": "本文から判別困難",
            }
        )
    return candidates[:3]


def candidate_label(candidate: Dict[str, str], index: int) -> str:
    page = candidate.get("page_number") or "?"
    label = candidate.get("label") or f"候補{index + 1}"
    caption = candidate.get("caption_or_content") or ""
    return f"候補{index + 1}: {label} / Page {page} / {caption[:45]}"


def apply_selected_figure_references(summary: Dict[str, Any], slide2_ref: Dict[str, str], slide3_ref: Dict[str, str]) -> Dict[str, Any]:
    updated = json.loads(json.dumps(summary, ensure_ascii=False))
    updated.setdefault("slide2", {})["figure_table_reference"] = slide2_ref
    updated.setdefault("slide3", {})["figure_table_reference"] = slide3_ref
    return updated


def render_pdf_page_png(pdf_bytes: bytes, page_number: int) -> Optional[tuple[io.BytesIO, int, int]]:
    try:
        import fitz
    except ImportError as error:
        raise RuntimeError("PDFページ画像の貼り付けには PyMuPDF が必要です。`pip install -r requirements.txt` を実行してください。") from error

    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    if len(document) == 0:
        return None
    page_index = max(0, min(page_number - 1, len(document) - 1))
    page = document.load_page(page_index)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
    pixel_width = pixmap.width
    pixel_height = pixmap.height
    stream = io.BytesIO(pixmap.tobytes("png"))
    stream.seek(0)
    document.close()
    return stream, pixel_width, pixel_height


def create_annotated_pdf(pdf_bytes: bytes, summary: Dict[str, Any]) -> bytes:
    try:
        import fitz
    except ImportError as error:
        raise RuntimeError("注釈PDFの作成には PyMuPDF が必要です。`pip install -r requirements.txt` を実行してください。") from error

    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    references = [
        ("Slide 2", "研究デザインを補完するFigure/Table", get_figure_table_reference(summary, "slide2")),
        ("Slide 3", "主要結果を補完するFigure/Table", get_figure_table_reference(summary, "slide3")),
    ]

    for slide_label, purpose, reference in references:
        page_number = parse_page_number(reference.get("page_number"))
        if not page_number or page_number > len(document):
            continue
        page = document.load_page(page_number - 1)
        rect = page.rect
        page.draw_rect(rect + (6, 6, -6, -6), color=(1, 0.78, 0), width=3)
        note_rect = fitz.Rect(rect.x0 + 24, rect.y0 + 24, min(rect.x1 - 24, rect.x0 + 430), rect.y0 + 112)
        note_text = (
            f"{slide_label}で引用候補として使用\n"
            f"{purpose}\n"
            f"{reference.get('label', 'Figure/Table')} / Page {reference.get('page_number') or '?'}\n"
            f"{reference.get('caption_or_content', '')[:120]}"
        )
        page.draw_rect(note_rect, color=(1, 0.72, 0), fill=(1, 0.96, 0.55), width=1)
        page.insert_textbox(note_rect + (8, 6, -8, -6), note_text, fontsize=8.5, color=(0, 0, 0), fontname="helv")
        page.add_text_annot(fitz.Point(note_rect.x1 + 8, note_rect.y0 + 8), note_text)

    output = io.BytesIO()
    document.save(output, garbage=4, deflate=True)
    document.close()
    output.seek(0)
    return output.getvalue()


def split_pdf_text_for_translation(pdf_text: str, max_chars: int = 9000) -> List[str]:
    page_blocks = re.split(r"(?=\[Page \d+\])", pdf_text)
    chunks: List[str] = []
    current = ""
    for block in page_blocks:
        block = block.strip()
        if not block:
            continue
        if current and len(current) + len(block) + 2 > max_chars:
            chunks.append(current.strip())
            current = block
        else:
            current = f"{current}\n\n{block}".strip() if current else block
    if current:
        chunks.append(current.strip())
    return chunks


def create_full_translation_markdown(client: OpenAI, model: str, pdf_text: str, summary: Dict[str, Any]) -> str:
    paper = summary.get("paper", {}) if isinstance(summary, dict) else {}
    chunks = split_pdf_text_for_translation(pdf_text)
    translated_chunks: List[str] = []
    system = "あなたは歯科医学論文を原文に忠実に日本語へ全文翻訳する専門家です。"

    for index, chunk in enumerate(chunks, start=1):
        prompt = f"""
以下のPDF抽出テキストを日本語に全文翻訳してください。

厳守事項:
- 要約しない。省略しない。
- [Page N] のページ表記を残す。
- 見出し、本文、図表キャプション、表中テキスト、統計値、P値、信頼区間、略語をできるだけ保持する。
- PDF抽出の崩れがある場合は、意味を補いすぎず自然な日本語に整える。
- 判別できない箇所は「判読困難」と書く。
- 出力はMarkdown本文のみ。

翻訳対象チャンク {index}/{len(chunks)}:
{chunk}
""".strip()
        translated = call_openai_text(client, model, system, prompt)
        translated_chunks.append(f"## 翻訳 Part {index}\n\n{translated.strip()}")

    header = [
        "# 論文全文日本語訳",
        "",
        f"**Title:** {paper.get('title_en') or '本文から判別困難'}",
        f"**Citation:** {paper.get('citation') or '本文から判別困難'}",
        f"**Article type:** {paper.get('article_type') or '本文から判別困難'}",
        "",
        "> 注意: PDFから抽出したテキストをもとにした全文翻訳です。PDF抽出の崩れがある箇所は、原著PDFで確認してください。",
        "",
    ]
    return "\n".join(header + translated_chunks)


def add_pdf_reference_image(
    slide,
    pdf_bytes: Optional[bytes],
    reference: Dict[str, str],
    left: float,
    top: float,
    width: float,
    height: float,
    title: str,
):
    add_textbox(slide, title, left, top - 0.34, width, 0.28, font_size=14, bold=True, color=RGBColor(0, 0, 0), align=PP_ALIGN.CENTER, font_name=JP_FONT_NAME, margin=0)
    page_number = parse_page_number(reference.get("page_number"))
    if pdf_bytes and page_number:
        rendered = render_pdf_page_png(pdf_bytes, page_number)
        if rendered:
            image_stream, pixel_width, pixel_height = rendered
            image_aspect = pixel_width / pixel_height
            box_aspect = width / height
            if image_aspect >= box_aspect:
                picture_width = width
                picture_height = width / image_aspect
            else:
                picture_height = height
                picture_width = height * image_aspect
            picture_left = left + (width - picture_width) / 2
            picture_top = top + (height - picture_height) / 2
            disable_shadow(slide.shapes.add_picture(image_stream, Inches(picture_left), Inches(picture_top), width=Inches(picture_width), height=Inches(picture_height)))
    else:
        placeholder = add_rounded_box(slide, left, top, width, height, RGBColor(248, 248, 248), RGBColor(150, 150, 150), 1.0)
        set_shape_text(
            placeholder,
            "図表ページを特定できませんでした。\n生成JSONの figure_table_reference を確認してください。",
            font_size=13,
            bold=True,
            align=PP_ALIGN.CENTER,
            font_name=JP_FONT_NAME,
        )
    add_rounded_box(slide, left, top, width, height, None, RGBColor(80, 110, 145), 1.0)
    caption = f"{reference.get('label', 'Figure/Table')} / Page {reference.get('page_number') or '?'}: {reference.get('caption_or_content', '')}"
    add_textbox(slide, caption[:130], left, top + height + 0.06, width, 0.26, font_size=10, bold=False, color=RGBColor(80, 80, 80), align=PP_ALIGN.CENTER, font_name=JP_FONT_NAME, margin=0)


def normalize_design_schema(summary: Dict[str, Any]) -> Dict[str, Any]:
    slide2 = summary.get("slide2", {})
    schema = slide2.get("design_schema") if isinstance(slide2, dict) else {}
    if not isinstance(schema, dict):
        schema = {}

    steps = as_list(schema.get("steps"))[:4]
    if not steps:
        steps = as_list(slide2.get("methods"))[:4]

    groups: List[Dict[str, str]] = []
    raw_groups = schema.get("groups")
    if isinstance(raw_groups, list):
        for item in raw_groups:
            if isinstance(item, dict):
                name = str(item.get("name", "")).strip()
                description = str(item.get("description", "")).strip()
            else:
                name = str(item).strip()
                description = ""
            if name or description:
                groups.append({"name": name or "比較群", "description": description})
    if not groups:
        for item in as_list(slide2.get("groups_or_comparison"))[:4]:
            groups.append({"name": item, "description": ""})

    labels = as_list(schema.get("labels"))[:4]
    return {
        "title": str(schema.get("title", "")).strip(),
        "steps": steps[:4],
        "groups": groups[:4],
        "labels": labels,
    }


def add_research_design_slide(slide, summary: Dict[str, Any], heading: str, pdf_bytes: Optional[bytes] = None):
    slide2 = summary.get("slide2", {})
    add_textbox(slide, heading, 3.5, 0.28, 3.0, 0.48, font_size=26, bold=True, color=RGBColor(0, 0, 0), align=PP_ALIGN.CENTER, font_name=JP_FONT_NAME, margin=0)

    top_items = []
    design_summary = get_nested(summary, "slide2", "design_summary", default="")
    participants = get_nested(summary, "slide2", "participants", default="")
    criteria = as_list(slide2.get("selection_criteria"))[:3]
    if design_summary:
        top_items.append(f"研究デザイン：{design_summary}")
    if participants:
        top_items.append(f"対象：{participants}")
    if criteria:
        top_items.append("選定条件：" + " / ".join(criteria))
    top_text = "\n".join(f"・{item}" for item in top_items[:4]) or "・本文から判別困難"
    top_box = add_rounded_box(slide, 0.55, 0.98, 8.9, 1.18, RGBColor(221, 235, 247), RGBColor(47, 78, 130), 1.0)
    set_shape_text(top_box, top_text, font_size=15, bold=True, font_name=JP_FONT_NAME, line_spacing=1.05)

    schema = normalize_design_schema(summary)
    steps = schema["steps"]
    if steps:
        step_count = min(len(steps), 4)
        box_w = 1.82 if step_count >= 4 else 2.2
        gap = 0.38
        start_x = max(0.75, (10 - (step_count * box_w + (step_count - 1) * gap)) / 2)
        y = 2.76
        for index, step in enumerate(steps[:4]):
            x = start_x + index * (box_w + gap)
            fill = [RGBColor(244, 224, 211), RGBColor(232, 243, 225), RGBColor(221, 235, 247), RGBColor(252, 240, 201)][index % 4]
            box = add_rounded_box(slide, x, y, box_w, 0.74, fill, RGBColor(170, 170, 170), 0.8)
            set_shape_text(box, step, font_size=12, bold=True, align=PP_ALIGN.CENTER, margin_left=0.08, margin_right=0.08, margin_top=0.06)
            if index < step_count - 1:
                add_arrow(slide, x + box_w + 0.05, y + 0.23, 0.28, 0.22)

    groups = schema["groups"]
    if groups:
        group_count = min(len(groups), 4)
        box_w = 1.85 if group_count >= 4 else 2.15
        gap = 0.38
        start_x = max(0.75, (10 - (group_count * box_w + (group_count - 1) * gap)) / 2)
        y = 3.78
        for index, group in enumerate(groups[:4]):
            x = start_x + index * (box_w + gap)
            fill = [RGBColor(242, 224, 212), RGBColor(250, 240, 203), RGBColor(221, 235, 247), RGBColor(232, 243, 225)][index % 4]
            text = group["name"]
            if group.get("description"):
                text = f"{text}\n{group['description']}"
            box = add_rounded_box(slide, x, y, box_w, 0.8, fill, RGBColor(190, 190, 190), 0.8)
            set_shape_text(box, text, font_size=12, bold=True, align=PP_ALIGN.CENTER, margin_left=0.08, margin_right=0.08, margin_top=0.06)

    outcomes = as_list(slide2.get("outcomes"))
    outcome_text = "\n".join(f"・{item}" for item in outcomes[:6]) or "・本文から判別困難"
    add_textbox(slide, "評価項目", 0.62, 5.22, 1.5, 0.34, font_size=19, bold=True, color=RGBColor(0, 0, 0), font_name=JP_FONT_NAME, margin=0)
    outcome_box = add_rounded_box(slide, 0.72, 5.62, 8.55, 1.12, QUESTION_FILL, RGBColor(0, 0, 0), 1.0)
    set_shape_text(outcome_box, outcome_text, font_size=15, bold=True, font_name=JP_FONT_NAME, line_spacing=1.05)


def normalize_evidence_items(summary: Dict[str, Any]) -> List[Dict[str, str]]:
    schema = summary.get("slide3", {}).get("evidence_schema")
    items: List[Dict[str, str]] = []
    if isinstance(schema, dict):
        raw_items = schema.get("items", [])
        if isinstance(raw_items, list):
            for item in raw_items:
                if not isinstance(item, dict):
                    continue
                label = str(item.get("label", "")).strip()
                value = str(item.get("value", "")).strip()
                interpretation = str(item.get("interpretation", "")).strip()
                if label or value or interpretation:
                    items.append(
                        {
                            "label": label or "評価項目",
                            "value": value or "本文から判別困難",
                            "interpretation": interpretation or "本文から判別困難",
                        }
                    )

    if items:
        return items[:4]

    key_numbers = as_list(summary.get("slide3", {}).get("key_numbers"))[:3]
    statistical_notes = as_list(summary.get("slide3", {}).get("statistical_notes"))[:3]
    main_results = as_list(summary.get("slide3", {}).get("main_results"))[:3]
    fallback_count = max(len(key_numbers), len(statistical_notes), len(main_results), 1)
    for index in range(min(fallback_count, 4)):
        items.append(
            {
                "label": f"根拠{index + 1}",
                "value": key_numbers[index] if index < len(key_numbers) else (main_results[index] if index < len(main_results) else "本文から判別困難"),
                "interpretation": statistical_notes[index] if index < len(statistical_notes) else (main_results[index] if index < len(main_results) else "本文から判別困難"),
            }
        )
    return items


def has_significance_phrase(text: str) -> bool:
    normalized = str(text).lower()
    patterns = [
        "有意",
        "p<",
        "p <",
        "p=",
        "p =",
        "p値",
        "信頼区間",
        "ci",
        "本文から判別困難",
    ]
    return any(pattern in normalized for pattern in patterns)


def format_statistical_note(note: str) -> str:
    text = str(note).strip()
    text = re.sub(r"^（(.+)）$", r"\1", text)
    text = text.replace("（", "、").replace("）", "")
    text = re.sub(r"\s+", " ", text).strip(" 、")
    return text


def main_results_with_significance(summary: Dict[str, Any]) -> List[str]:
    results = as_list(summary.get("slide3", {}).get("main_results"))[:3] or ["本文から判別困難"]
    notes = as_list(summary.get("slide3", {}).get("statistical_notes"))
    adjusted: List[str] = []
    for index, result in enumerate(results):
        text = str(result).strip()
        if has_significance_phrase(text):
            adjusted.append(text)
            continue
        note = notes[index] if index < len(notes) else (notes[0] if len(notes) == 1 else "")
        if note and str(note).strip() and "本文から判別困難" not in str(note):
            adjusted.append(f"{text}（{format_statistical_note(note)}）")
        else:
            adjusted.append(f"{text}（有意差は本文から判別困難）")
    return adjusted


def estimate_result_box_layout(items: List[str]) -> Dict[str, float]:
    total_lines = 0
    for item in items:
        total_lines += max(1, (len(str(item)) + 34) // 35)
    font_size = 17
    if total_lines >= 6:
        font_size = 14
    elif total_lines >= 5:
        font_size = 15
    height = min(1.65, max(1.15, 0.28 + total_lines * 0.28))
    return {"height": height, "font_size": font_size}


def add_evidence_schema(slide, summary: Dict[str, Any], top: float = 2.68):
    schema = summary.get("slide3", {}).get("evidence_schema")
    title = schema.get("title") if isinstance(schema, dict) else ""
    title = str(title).strip() or "主要結果の根拠"
    add_textbox(slide, title, 2.75, top - 0.5, 4.5, 0.34, font_size=20, bold=False, color=RGBColor(0, 0, 0), align=PP_ALIGN.CENTER, font_name=JP_FONT_NAME, margin=0)

    items = normalize_evidence_items(summary)
    left = 0.78
    width = 8.44
    row_h = 0.48
    col_w = [2.1, 2.8, 3.54]
    headers = ["評価・比較", "数値/所見", "読み取り"]
    x = left
    for index, header in enumerate(headers):
        cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(top), Inches(col_w[index]), Inches(row_h))
        disable_shadow(cell)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(88, 126, 164)
        cell.line.color.rgb = RGBColor(255, 255, 255)
        add_textbox(slide, header, x + 0.05, top + 0.1, col_w[index] - 0.1, 0.25, font_size=12, bold=True, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER, font_name=JP_FONT_NAME, margin=0)
        x += col_w[index]

    for row_index, item in enumerate(items[:4], start=1):
        y = top + row_h * row_index
        x = left
        fill = RGBColor(235, 242, 250) if row_index % 2 else RGBColor(248, 250, 252)
        values = [item["label"], item["value"], item["interpretation"]]
        for col_index, value in enumerate(values):
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(col_w[col_index]), Inches(row_h))
            disable_shadow(cell)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            cell.line.color.rgb = RGBColor(255, 255, 255)
            add_textbox(
                slide,
                value,
                x + 0.08,
                y + 0.06,
                col_w[col_index] - 0.16,
                0.34,
                font_size=10,
                bold=(col_index == 0),
                color=TEXT_COLOR,
                align=PP_ALIGN.CENTER if col_index == 0 else PP_ALIGN.LEFT,
                font_name=JP_FONT_NAME,
                margin=0,
            )
            x += col_w[col_index]

    add_rounded_box(slide, left - 0.03, top - 0.03, width + 0.06, row_h * (len(items[:4]) + 1) + 0.06, None, RGBColor(88, 126, 164), 1.0)


def add_research_results_slide(slide, summary: Dict[str, Any], heading: str, pdf_bytes: Optional[bytes] = None):
    add_textbox(slide, heading, 4.1, 0.24, 1.8, 0.44, font_size=24, bold=True, color=RGBColor(0, 0, 0), align=PP_ALIGN.CENTER, font_name=JP_FONT_NAME, margin=0)

    main_results = main_results_with_significance(summary)
    result_layout = estimate_result_box_layout(main_results)
    result_box_height = result_layout["height"]
    result_font_size = int(result_layout["font_size"])
    top_box = add_rounded_box(slide, 0.55, 0.84, 8.9, result_box_height, RGBColor(221, 235, 247), RGBColor(65, 82, 105), 1.0)
    top_box.text_frame.clear()
    top_box.text_frame.margin_left = Inches(0.18)
    top_box.text_frame.margin_right = Inches(0.18)
    top_box.text_frame.margin_top = Inches(0.08)
    top_box.text_frame.margin_bottom = Inches(0.04)
    top_box.text_frame.word_wrap = True
    for index, item in enumerate(main_results):
        paragraph = top_box.text_frame.paragraphs[0] if index == 0 else top_box.text_frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 1.05
        run = paragraph.add_run()
        run.text = f"・{item}"
        run.font.name = JP_FONT_NAME
        run.font.size = Pt(result_font_size)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    evidence_top = 0.84 + result_box_height + 0.55
    add_evidence_schema(slide, summary, top=evidence_top)

    detail_left = 0.55
    detail_top = 5.7
    detail_width = 8.9
    detail_height = 0.82
    add_rounded_box(slide, detail_left, detail_top, detail_width, detail_height, None, RGBColor(206, 58, 105), 1.0)
    add_textbox(
        slide,
        "さらに詳しく",
        0.18,
        detail_top - 0.04,
        1.35,
        0.34,
        font_size=17,
        bold=True,
        color=RGBColor(0, 0, 0),
        font_name=JP_FONT_NAME,
        rotation=340,
        margin=0,
        word_wrap=False,
        fill_color=RGBColor(255, 255, 255),
    )
    add_multiline_text(
        slide,
        get_nested(summary, "slide3", "more_detail"),
        detail_left + 0.55,
        detail_top + 0.22,
        detail_width - 1.0,
        detail_height - 0.32,
        font_size=14,
        bold=False,
        line_spacing=1.1,
        font_name=JP_FONT_NAME,
    )


def circled_number(index: int) -> str:
    numbers = ["①", "②", "③", "④", "⑤"]
    return numbers[index] if index < len(numbers) else f"{index + 1})"


def create_pptx(summary: Dict[str, Any], box_number: str = "000", pdf_bytes: Optional[bytes] = None) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]
    review_article = is_review_article(summary)
    slide2_heading = "本レビューの概要①" if review_article else "研究デザイン"
    slide3_heading = "本レビューの概要②" if review_article else "結果"

    slide = prs.slides.add_slide(blank)
    add_slide_background(slide)
    add_title_header(slide, box_number, get_nested(summary, "paper", "title_en"), get_nested(summary, "paper", "citation"))
    question = get_nested(summary, "slide1", "clinical_question_jp")
    key_message = get_nested(summary, "slide1", "key_message_jp")
    add_pill(
        slide,
        question,
        0.7,
        1.94,
        8.6,
        0.74,
        QUESTION_FILL,
        font_size=font_size_for(question, 20, 18, 34),
        bold=True,
    )
    if CQ_TRIANGLE_PATH.exists():
        disable_shadow(slide.shapes.add_picture(str(CQ_TRIANGLE_PATH), Inches(1.35), Inches(2.84), height=Inches(0.78)))
    add_textbox(slide, key_message, 2.1, 2.95, 6.85, 0.6, font_size=font_size_for(key_message, 18, 16, 42), bold=True, color=RGBColor(0, 0, 0), margin=0)
    if CQ_UNDERLINE_PATH.exists():
        disable_shadow(slide.shapes.add_picture(str(CQ_UNDERLINE_PATH), Inches(1.78), Inches(3.63), width=Inches(6.95), height=Inches(0.06)))
    add_section_heading(slide, get_nested(summary, "slide1", "background_heading"), 0.42, 3.82, 2.6)
    add_multiline_text(slide, get_nested(summary, "slide1", "background_overview"), 0.9, 4.18, 8.35, 3.05, font_size=15)

    slide = prs.slides.add_slide(blank)
    add_slide_background(slide)
    if review_article:
        add_textbox(slide, slide2_heading, 3.65, 0.35, 2.7, 0.45, font_size=24, bold=True, color=TEXT_COLOR, margin=0)
        add_pill(slide, get_nested(summary, "slide2", "design_summary"), 0.55, 0.95, 5.9, 0.45, QUESTION_FILL, font_size=17, bold=True)
        add_result_card(slide, "対象", get_nested(summary, "slide2", "participants"), 0.55, 1.65, 4.2, 1.2)
        add_result_card(slide, "評価項目", "\n".join(as_list(summary.get("slide2", {}).get("outcomes"))), 5.15, 1.65, 4.1, 1.2)
        add_section_heading(slide, "方法", 0.7, 3.25, 1.4)
        add_bullet_lines(slide, as_list(summary.get("slide2", {}).get("methods")), 0.95, 3.68, 4.0, 2.4, font_size=15)
        add_section_heading(slide, "比較・群分け", 5.25, 3.25, 2.0)
        add_bullet_lines(slide, as_list(summary.get("slide2", {}).get("groups_or_comparison")), 5.45, 3.68, 3.75, 2.4, font_size=15)
    else:
        add_research_design_slide(slide, summary, slide2_heading, pdf_bytes=pdf_bytes)

    slide = prs.slides.add_slide(blank)
    add_slide_background(slide)
    if review_article:
        add_textbox(slide, slide3_heading, 3.65, 0.32, 2.9, 0.45, font_size=24, bold=True, color=TEXT_COLOR, align=PP_ALIGN.CENTER, font_name=JP_FONT_NAME, margin=0)
        add_result_card(slide, "主な結果", "\n".join(f"・{item}" for item in as_list(summary.get("slide3", {}).get("main_results"))), 0.62, 1.02, 8.65, 1.55)
        add_result_card(slide, "重要な数値", "\n".join(as_list(summary.get("slide3", {}).get("key_numbers"))), 0.72, 2.92, 4.05, 1.55)
        add_result_card(slide, "統計", "\n".join(as_list(summary.get("slide3", {}).get("statistical_notes"))), 5.18, 2.92, 4.0, 1.55)
        add_section_heading(slide, "さらに詳しく", 0.62, 5.58, 1.7)
        add_multiline_text(slide, get_nested(summary, "slide3", "more_detail"), 1.65, 5.95, 7.45, 0.9, font_size=15)
    else:
        add_research_results_slide(slide, summary, slide3_heading, pdf_bytes=pdf_bytes)

    slide = prs.slides.add_slide(blank)
    add_slide_background(slide)
    selection_title_rotation = 350
    add_textbox(
        slide,
        "論文選定のポイント",
        0.55,
        0.5,
        3.2,
        0.35,
        font_size=20,
        bold=True,
        color=RGBColor(0, 0, 0),
        font_name=JP_FONT_NAME,
        rotation=selection_title_rotation,
        margin=0,
        word_wrap=False,
    )
    underline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(0.86), Inches(2.45), Inches(0.025))
    disable_shadow(underline)
    underline.fill.solid()
    underline.fill.fore_color.rgb = RGBColor(0, 160, 235)
    underline.line.color.rgb = RGBColor(0, 160, 235)
    underline.rotation = selection_title_rotation
    points = as_list(summary.get("slide4", {}).get("selection_points"))[:2]
    point_text = "\n".join(f"{circled_number(index)}{point}" for index, point in enumerate(points or ["本文から判別困難"]))
    add_multiline_text(slide, point_text, 1.35, 1.12, 7.65, 0.9, font_size=18, bold=False, line_spacing=1.05, font_name=JP_FONT_NAME)

    if DOCTOR_COMMENT_ICON_PATH.exists():
        disable_shadow(slide.shapes.add_picture(str(DOCTOR_COMMENT_ICON_PATH), Inches(0.28), Inches(1.95), width=Inches(0.34), height=Inches(0.34)))
        disable_shadow(slide.shapes.add_picture(str(DOCTOR_COMMENT_ICON_PATH), Inches(3.8), Inches(1.95), width=Inches(0.34), height=Inches(0.34)))
    add_textbox(
        slide,
        "抄読ドクターからひとこと",
        0.68,
        1.93,
        3.1,
        0.34,
        font_size=20,
        bold=True,
        color=RGBColor(0, 0, 0),
        font_name=JP_FONT_NAME,
        margin=0,
        word_wrap=False,
    )

    doctor_bg_left = 0.55
    doctor_bg_top = 2.22
    doctor_bg_width = 8.95
    doctor_bg_height = 4.55
    doctor_text_margin_x = 0.4
    doctor_text_margin_y = 0.4
    if DOCTOR_COMMENT_BG_PATH.exists():
        disable_shadow(slide.shapes.add_picture(str(DOCTOR_COMMENT_BG_PATH), Inches(doctor_bg_left), Inches(doctor_bg_top), width=Inches(doctor_bg_width), height=Inches(doctor_bg_height)))
    add_multiline_text(
        slide,
        get_nested(summary, "slide4", "doctor_comment"),
        doctor_bg_left + doctor_text_margin_x,
        doctor_bg_top + doctor_text_margin_y,
        doctor_bg_width - doctor_text_margin_x * 2,
        doctor_bg_height - doctor_text_margin_y * 2,
        font_size=18,
        bold=False,
        line_spacing=1.15,
        font_name=JP_FONT_NAME,
    )

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def safe_filename(name: str) -> str:
    stem = Path(name).stem
    stem = re.sub(r"[^\w\-一-龥ぁ-んァ-ンー]+", "_", stem).strip("_")
    return stem or "periobox_summary"


def main():
    st.set_page_config(page_title=APP_TITLE, layout="wide")
    current_user = require_authorized_user()
    if current_user is None:
        return

    st.title(APP_TITLE)
    st.caption("歯科医療論文PDFから、編集しやすい4枚構成の学術スライドを生成します。")

    with st.sidebar:
        st.header("設定")
        server_api_key_set = bool(get_secret_value("OPENAI_API_KEY"))
        if server_api_key_set:
            api_key = ""
            st.success("OpenAI APIキーはサーバー側で設定済みです。")
        else:
            api_key = st.text_input(
                "OpenAI APIキー",
                value="",
                type="password",
                help="ローカル検証用です。本番運用ではサーバー側のOPENAI_API_KEYに設定してください。",
            )
        model = st.text_input("OpenAI model", value=DEFAULT_MODEL)
        box_number = st.text_input("PerioBox番号", value="000", help="左上に表示する番号です。不要なら000のままで大丈夫です。")
        st.caption("本番ではOPENAI_API_KEYをSecrets/環境変数に保存し、モデレーターには表示しません。")

    uploaded_pdf = st.file_uploader("論文PDFをアップロード", type=["pdf"])

    if uploaded_pdf:
        if st.button("PowerPointを生成", type="primary"):
            try:
                with st.status("PDF本文を抽出しています...", expanded=True) as status:
                    pdf_bytes = uploaded_pdf.getvalue()
                    pdf_text = extract_pdf_text(io.BytesIO(pdf_bytes))
                    if not pdf_text.strip():
                        raise RuntimeError("PDF本文を抽出できませんでした。スキャンPDFの場合はOCR済みPDFを使用してください。")

                    st.write(f"抽出文字数: {len(pdf_text):,} 文字")
                    client = get_openai_client(api_key)

                    status.update(label="PDF本文をセクション別に抽出しています...")
                    sections = extract_sections_with_openai(client, model, pdf_text)

                    status.update(label="原文に忠実な日本語構造化要約を生成しています...")
                    faithful_summary = create_faithful_summary_with_openai(client, model, sections)

                    status.update(label="PerioBoxスライド用要約へ変換しています...")
                    summary = create_slide_summary_with_openai(client, model, faithful_summary)

                    status.update(label="PowerPointを作成しています...")
                    pptx_bytes = create_pptx(summary, box_number=box_number, pdf_bytes=pdf_bytes)
                    status.update(label="論文全文の日本語訳Markdownを生成しています...")
                    translation_markdown = create_full_translation_markdown(client, model, pdf_text, summary)
                    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                    output_stem = f"{safe_filename(uploaded_pdf.name)}_{timestamp}"
                    output_name = f"{output_stem}.pptx"
                    translation_name = f"{output_stem}_translation.md"
                    OUTPUT_DIR.mkdir(exist_ok=True)
                    output_path = OUTPUT_DIR / output_name
                    translation_path = OUTPUT_DIR / translation_name
                    output_path.write_bytes(pptx_bytes)
                    translation_path.write_text(translation_markdown, encoding="utf-8")

                    st.session_state.pop("pptx_bytes", None)
                    st.session_state.pop("pptx_file_name", None)
                    st.session_state.pop("pptx_output_path", None)
                    st.session_state.pop("annotated_pdf_bytes", None)
                    st.session_state.pop("annotated_pdf_file_name", None)
                    st.session_state.pop("annotated_pdf_output_path", None)
                    st.session_state["summary"] = summary
                    st.session_state["sections"] = sections
                    st.session_state["faithful_summary"] = faithful_summary
                    st.session_state["pdf_bytes"] = pdf_bytes
                    st.session_state["uploaded_pdf_name"] = uploaded_pdf.name
                    st.session_state["pptx_bytes"] = pptx_bytes
                    st.session_state["pptx_file_name"] = output_name
                    st.session_state["pptx_output_path"] = str(output_path.resolve())
                    st.session_state["translation_markdown"] = translation_markdown
                    st.session_state["translation_file_name"] = translation_name
                    st.session_state["translation_output_path"] = str(translation_path.resolve())
                    status.update(label="完了しました。", state="complete")

                st.success("4枚構成PowerPointを生成しました。")
                st.info(f"ファイルも保存しました: {st.session_state['pptx_output_path']}")
                st.info(f"日本語訳Markdownも保存しました: {st.session_state['translation_output_path']}")
            except Exception as error:
                st.error(str(error))

    if "pptx_bytes" in st.session_state:
        st.download_button(
            label="PowerPointをダウンロード",
            data=st.session_state["pptx_bytes"],
            file_name=st.session_state.get("pptx_file_name", "periobox_style_summary.pptx"),
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            key="download_pptx",
        )

        if "pptx_output_path" in st.session_state:
            st.caption(f"保存先: {st.session_state['pptx_output_path']}")

        if "annotated_pdf_bytes" in st.session_state:
            st.download_button(
                label="引用注釈つきPDFをダウンロード",
                data=st.session_state["annotated_pdf_bytes"],
                file_name=st.session_state.get("annotated_pdf_file_name", "annotated_paper.pdf"),
                mime="application/pdf",
                key="download_annotated_pdf",
            )
            st.caption(f"保存先: {st.session_state.get('annotated_pdf_output_path')}")

        if "translation_markdown" in st.session_state:
            st.download_button(
                label="日本語訳Markdownをダウンロード",
                data=st.session_state["translation_markdown"].encode("utf-8"),
                file_name=st.session_state.get("translation_file_name", "paper_translation.md"),
                mime="text/markdown",
                key="download_translation_md",
            )
            st.caption(f"保存先: {st.session_state.get('translation_output_path')}")

        if "summary" in st.session_state:
            with st.expander("生成された構造化要約JSON"):
                st.json(st.session_state["summary"])

        if "faithful_summary" in st.session_state:
            with st.expander("忠実な日本語構造化要約JSON"):
                st.json(st.session_state["faithful_summary"])

        if "sections" in st.session_state:
            with st.expander("セクション別抽出JSON"):
                st.json(st.session_state["sections"])


if __name__ == "__main__":
    main()

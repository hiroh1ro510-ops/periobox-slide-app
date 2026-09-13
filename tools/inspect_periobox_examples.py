import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor


EMU_PER_INCH = 914400


DESIGN_TITLE_PATTERNS = [
    r"研\s*究\s*デ\s*ザ\s*イ\s*ン",
    r"研\s*究\s*方\s*法",
    r"方\s*法",
    r"デ\s*ザ\s*イ\s*ン",
    r"study\s*design",
    r"methods?",
    r"materials?\s*(and|&)\s*methods?",
]

RESULT_TITLE_PATTERNS = [
    r"結\s*果",
    r"主\s*要\s*結\s*果",
    r"results?",
    r"outcomes?",
    r"findings?",
]


def emu_to_in(value):
    if value is None:
        return None
    return round(value / EMU_PER_INCH, 3)


def rgb_to_hex(rgb):
    if rgb is None:
        return None
    try:
        return f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    except Exception:
        return None


def safe_rgb(color_format):
    try:
        return rgb_to_hex(color_format.rgb)
    except Exception:
        return None


def fill_info(shape):
    info = {"type": None, "fore_rgb": None}
    try:
        fill = shape.fill
    except Exception:
        return info
    try:
        info["type"] = str(fill.type)
    except Exception:
        pass
    try:
        info["fore_rgb"] = safe_rgb(fill.fore_color)
    except Exception:
        pass
    return info


def line_info(shape):
    info = {"rgb": None, "width_pt": None}
    try:
        line = shape.line
    except Exception:
        return info
    try:
        info["rgb"] = safe_rgb(line.color)
    except Exception:
        pass
    try:
        if line.width is not None:
            info["width_pt"] = round(line.width.pt, 2)
    except Exception:
        pass
    return info


def run_text_info(shape):
    infos = []
    if not getattr(shape, "has_text_frame", False):
        return infos
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            text = run.text
            if not text or not text.strip():
                continue
            color = None
            try:
                color = rgb_to_hex(run.font.color.rgb)
            except Exception:
                color = None
            infos.append(
                {
                    "text": text,
                    "font_size": round(run.font.size.pt, 1) if run.font.size else None,
                    "font_name": run.font.name,
                    "bold": run.font.bold,
                    "italic": run.font.italic,
                    "color": color,
                }
            )
    return infos


def shape_text(shape):
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()


def autoshape_name(shape):
    try:
        return str(shape.auto_shape_type) if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE else None
    except Exception:
        return None


def picture_info(shape):
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        return None
    info = {"filename": None, "content_type": None, "size_bytes": None}
    try:
        image = shape.image
        info["filename"] = image.filename
        info["content_type"] = image.content_type
        info["size_bytes"] = len(image.blob)
    except Exception:
        pass
    return info


def group_shapes(shape):
    """Flatten group shapes into a list of leaf shapes."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        out = []
        for child in shape.shapes:
            out.extend(group_shapes(child))
        return out
    return [shape]


def describe_shape(shape):
    item = {
        "shape_type": str(shape.shape_type),
        "auto_shape": autoshape_name(shape),
        "name": shape.name,
        "left": emu_to_in(shape.left),
        "top": emu_to_in(shape.top),
        "width": emu_to_in(shape.width),
        "height": emu_to_in(shape.height),
        "rotation": round(shape.rotation, 2) if getattr(shape, "rotation", None) else 0.0,
        "text": shape_text(shape),
        "runs": run_text_info(shape),
        "fill": fill_info(shape),
        "line": line_info(shape),
        "picture": picture_info(shape),
    }
    return item


def classify_slide(items):
    """Return a list of tags matched against title-like text on the slide."""
    tags = set()
    for item in items:
        text = (item.get("text") or "").strip()
        if not text:
            continue
        # Only consider short "title-ish" text
        if len(text) > 40:
            continue
        low = text.lower()
        for pattern in DESIGN_TITLE_PATTERNS:
            if re.search(pattern, text) or re.search(pattern, low):
                tags.add("design")
                break
        for pattern in RESULT_TITLE_PATTERNS:
            if re.search(pattern, text) or re.search(pattern, low):
                tags.add("result")
                break
    return sorted(tags)


def inspect_deck(path):
    prs = Presentation(path)
    deck = {
        "file": str(path),
        "slides": len(prs.slides),
        "width_in": emu_to_in(prs.slide_width),
        "height_in": emu_to_in(prs.slide_height),
        "slide_items": [],
    }
    for slide_index, slide in enumerate(prs.slides, start=1):
        if slide_index > 3:
            break  # Slides 4+ are known-garbage in these samples
        items = []
        for shape in slide.shapes:
            for leaf in group_shapes(shape):
                items.append(describe_shape(leaf))
        tags = classify_slide(items)
        deck["slide_items"].append({"slide": slide_index, "tags": tags, "items": items})
    return deck


def pick_design_and_result_slides(deck):
    """Slide 2 is design, slide 3 is result. Slide 4+ is ignored (garbage)."""
    design_slide = deck["slide_items"][1]["slide"] if len(deck["slide_items"]) >= 2 else None
    result_slide = deck["slide_items"][2]["slide"] if len(deck["slide_items"]) >= 3 else None
    return design_slide, result_slide


def summarize(decks):
    font_sizes = Counter()
    font_names = Counter()
    text_colors = Counter()
    fill_colors = Counter()
    line_colors = Counter()
    slide_counts = Counter()
    autoshapes = Counter()
    shape_types = Counter()
    picture_count = 0
    design_slides = []
    result_slides = []

    for deck in decks:
        slide_counts[deck["slides"]] += 1
        design_idx, result_idx = pick_design_and_result_slides(deck)
        design_slides.append({"file": deck["file"], "slide": design_idx})
        result_slides.append({"file": deck["file"], "slide": result_idx})
        for slide in deck["slide_items"]:
            for item in slide["items"]:
                shape_types[item["shape_type"]] += 1
                if item.get("auto_shape"):
                    autoshapes[item["auto_shape"]] += 1
                if item.get("picture"):
                    picture_count += 1
                if item["fill"].get("fore_rgb"):
                    fill_colors[item["fill"]["fore_rgb"]] += 1
                if item["line"].get("rgb"):
                    line_colors[item["line"]["rgb"]] += 1
                for run in item["runs"]:
                    if run["font_size"]:
                        font_sizes[str(run["font_size"])] += 1
                    if run["font_name"]:
                        font_names[run["font_name"]] += 1
                    if run["color"]:
                        text_colors[run["color"]] += 1

    return {
        "slide_counts": slide_counts.most_common(),
        "font_sizes": font_sizes.most_common(20),
        "font_names": font_names.most_common(20),
        "text_colors": text_colors.most_common(20),
        "fill_colors": fill_colors.most_common(20),
        "line_colors": line_colors.most_common(20),
        "autoshapes": autoshapes.most_common(30),
        "shape_types": shape_types.most_common(),
        "picture_count": picture_count,
        "design_slide_pick": design_slides,
        "result_slide_pick": result_slides,
    }


def extract_focus(decks):
    """Extract only the design and result slides for each deck."""
    focus = []
    for deck in decks:
        design_idx, result_idx = pick_design_and_result_slides(deck)
        entry = {"file": deck["file"], "design_slide": design_idx, "result_slide": result_idx, "slides": {}}
        for slide in deck["slide_items"]:
            if slide["slide"] in (design_idx, result_idx):
                role = "design" if slide["slide"] == design_idx else "result"
                entry["slides"][role] = slide
        focus.append(entry)
    return focus


def main():
    root = Path(sys.argv[1])
    out_dir = root
    decks = [inspect_deck(path) for path in sorted(root.glob("case*_periobox.pptx"))]
    payload = {
        "summary": summarize(decks),
        "decks": decks,
    }
    (out_dir / "_measurements.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    focus_payload = {
        "summary": payload["summary"],
        "focus": extract_focus(decks),
    }
    (out_dir / "_focus_design_results.json").write_text(
        json.dumps(focus_payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("wrote:")
    print(" ", out_dir / "_measurements.json")
    print(" ", out_dir / "_focus_design_results.json")


if __name__ == "__main__":
    main()
